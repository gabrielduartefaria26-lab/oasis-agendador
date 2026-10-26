import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from dados_ligar import LIGAR
from dados_concorrentes import CONCORRENTES, OPORTUNIDADES

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Lista_Prospeccao_Moagem_Boituva.xlsx")
F = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(name=F, bold=True, color="FFFFFF", size=10)
BODY = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
TITLE = Font(name=F, size=14, bold=True, color="1F4E78")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
PRIO_FILL = {"REF": "D9D9D9", "A": "C6EFCE", "B": "DDEBF7", "C": "FCE4D6"}
STATUS = ["A ligar", "Sem resposta", "Retornar", "Interessado", "Enviar proposta", "Fechado", "Sem interesse"]

wb = Workbook()

# ---------------- Como usar ----------------
ws = wb.active
ws.title = "Como usar"
ws["A1"] = "Prospecção: serviço de moagem de aparas plásticas (Boituva-SP)"
ws["A1"].font = TITLE
rows = [
    ("Abas", ""),
    ("Ligar", "Empresas para contato (clientes em potencial). Preencha as colunas amarelas a cada ligação."),
    ("Concorrentes", "Quem já presta serviço parecido e como vende."),
    ("Oportunidades", "O que os concorrentes fazem bem e as brechas para vocês."),
    ("Roteiro", "O que falar e o que perguntar na ligação."),
    ("", ""),
    ("Prioridade", ""),
    ("REF", "Empresas que você citou (Arqplast, MM Plást): servem de modelo do cliente ideal."),
    ("A", "Recicladora/granuladora de PP, PE ou PET: mesmo perfil do cliente ideal. Ligar primeiro."),
    ("B", "Fábrica que GERA apara de fita, ráfia, filme ou fio: pode pagar pela moagem ou vender a apara."),
    ("C", "Aparista ou gestora de resíduos: junta volume e pode terceirizar a moagem."),
    ("", ""),
    ("Atenção", "Telefones vieram de sites e diretórios públicos em out/2026. 'não verificado' = não confirmado; procure no Google Maps."),
    ("", "Distância (km) é estimativa aproximada por estrada a partir de Boituva."),
    ("", "Capacidade da máquina: até 4 t/dia (~88 t/mês em 22 dias úteis), segundo informação do usuário."),
]
r = 3
for a, b in rows:
    ws.cell(r, 1, a).font = BOLD
    ws.cell(r, 2, b).font = BODY
    ws.cell(r, 2).alignment = WRAP
    if a in PRIO_FILL:
        ws.cell(r, 1).fill = PatternFill("solid", fgColor=PRIO_FILL[a])
    r += 1

r += 1
ws.cell(r, 1, "Resumo das ligações").font = TITLE
r += 1
ws.cell(r, 1, "Status").font = HEAD_FONT; ws.cell(r, 1).fill = HEAD_FILL
ws.cell(r, 2, "Quantidade").font = HEAD_FONT; ws.cell(r, 2).fill = HEAD_FILL
r += 1
n_last = 1 + len(LIGAR)
first_sum = r
for s in STATUS:
    ws.cell(r, 1, s).font = BODY
    ws.cell(r, 2, f'=COUNTIF(Ligar!$N$2:$N${n_last},A{r})').font = BODY
    r += 1
ws.cell(r, 1, "Total de empresas").font = BOLD
ws.cell(r, 2, f"=COUNTA(Ligar!$C$2:$C${n_last})").font = BOLD
r += 1
ws.cell(r, 1, "Já contatadas").font = BOLD
ws.cell(r, 2, f'=B{r-1}-B{first_sum}').font = BOLD
ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 110

# ---------------- Ligar ----------------
ws = wb.create_sheet("Ligar")
headers = ["Prior.", "Nº", "Empresa", "Cidade", "Dist. aprox. (km)", "Tipo", "O que fazem", "Telefone",
           "WhatsApp", "E-mail / Site", "Confiança dos dados", "Por que ligar / o que oferecer", "Fonte",
           "Status", "Data do contato", "Falou com (nome/cargo)", "Volume de apara/mês", "Próximo passo", "Observações"]
widths = [7, 5, 30, 18, 9, 24, 42, 22, 24, 34, 11, 46, 28, 15, 12, 22, 14, 26, 30]
for i, (h, w) in enumerate(zip(headers, widths), 1):
    c = ws.cell(1, i, h)
    c.font = HEAD_FONT; c.fill = HEAD_FILL; c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BORDER
    ws.column_dimensions[c.column_letter].width = w
order = {"REF": 0, "A": 1, "B": 2, "C": 3}
data = sorted(LIGAR, key=lambda x: (order[x[0]], x[3]))
for idx, row in enumerate(data, 1):
    prio, emp, cid, km, tipo, oque, tel, wpp, site, conf, porque, fonte = row
    vals = [prio, idx, emp, cid, km, tipo, oque, tel, wpp, site, conf, porque, fonte,
            "A ligar", None, None, None, None, None]
    rr = idx + 1
    for i, v in enumerate(vals, 1):
        c = ws.cell(rr, i, v)
        c.font = BODY; c.alignment = WRAP; c.border = BORDER
        if i >= 14:
            c.fill = INPUT_FILL
    ws.cell(rr, 1).fill = PatternFill("solid", fgColor=PRIO_FILL[prio])
    ws.cell(rr, 1).font = BOLD
    ws.cell(rr, 3).font = BOLD
    ws.cell(rr, 15).number_format = "DD/MM/YYYY"
    if fonte.startswith("http"):
        ws.cell(rr, 13).hyperlink = fonte
        ws.cell(rr, 13).font = Font(name=F, size=10, color="0563C1", underline="single")
last = len(data) + 1
dv = DataValidation(type="list", formula1='"' + ",".join(STATUS) + '"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"N2:N{last}")
for s, color in [("Interessado", "C6EFCE"), ("Enviar proposta", "C6EFCE"), ("Fechado", "00B050"),
                 ("Sem interesse", "D9D9D9"), ("Retornar", "FFEB9C"), ("Sem resposta", "FFEB9C")]:
    ws.conditional_formatting.add(f"N2:N{last}",
        FormulaRule(formula=[f'$N2="{s}"'], fill=PatternFill("solid", fgColor=color)))
ws.freeze_panes = "D2"
ws.auto_filter.ref = f"A1:{ws.cell(1, len(headers)).column_letter}{last}"
ws.row_dimensions[1].height = 32

# ---------------- Concorrentes ----------------
ws = wb.create_sheet("Concorrentes")
headers = ["Região", "Empresa", "Cidade", "Dist. aprox. (km)", "Perfil", "Serviços", "Materiais / capacidade",
           "Telefone", "Site / redes", "Como vendem", "Ponto forte", "Ponto fraco / brecha", "Fonte"]
widths = [16, 26, 16, 9, 26, 44, 24, 26, 28, 40, 26, 30, 30]
for i, (h, w) in enumerate(zip(headers, widths), 1):
    c = ws.cell(1, i, h)
    c.font = HEAD_FONT; c.fill = HEAD_FILL; c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BORDER
    ws.column_dimensions[c.column_letter].width = w
for rr, row in enumerate(CONCORRENTES, 2):
    for i, v in enumerate(row, 1):
        c = ws.cell(rr, i, v)
        c.font = BODY; c.alignment = WRAP; c.border = BORDER
    ws.cell(rr, 2).font = BOLD
    if row[0].startswith("Refer"):
        ws.cell(rr, 1).fill = PatternFill("solid", fgColor="EDEDED")
    if row[-1].startswith("http"):
        ws.cell(rr, 13).hyperlink = row[-1]
        ws.cell(rr, 13).font = Font(name=F, size=10, color="0563C1", underline="single")
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:M{len(CONCORRENTES)+1}"
ws.row_dimensions[1].height = 32

# ---------------- Oportunidades ----------------
ws = wb.create_sheet("Oportunidades")
ws["A1"] = "Concorrência: o que aprender e onde se diferenciar"
ws["A1"].font = TITLE
r = 3
for titulo, itens in OPORTUNIDADES:
    ws.cell(r, 1, titulo).font = Font(name=F, size=11, bold=True, color="FFFFFF")
    ws.cell(r, 1).fill = HEAD_FILL
    r += 1
    for it in itens:
        c = ws.cell(r, 1, "• " + it); c.font = BODY; c.alignment = WRAP
        r += 1
    r += 1
ws.column_dimensions["A"].width = 130

# ---------------- Roteiro ----------------
ws = wb.create_sheet("Roteiro")
ws["A1"] = "Roteiro de ligação"
ws["A1"].font = TITLE
roteiro = [
    ("1. Abertura (30 s)", [
        "Bom dia, meu nome é ___, sou da ___, de Boituva. Com quem eu falo sobre aparas, resíduos ou compra de material reciclado?",
        "Somos especializados em moagem de fita de arquear, ráfia e plástico flexível (PP, PE e PET).",
        "Moemos até 4 toneladas por dia, com granulometria de até 2 mm, e entregamos em big bag.",
    ]),
    ("2. Perguntas (descobrir a dor)", [
        "Vocês geram ou compram apara de fita, ráfia ou filme? Mais ou menos quantas toneladas por mês?",
        "Hoje quem mói esse material? Fazem internamente ou terceirizam?",
        "A moagem interna está no limite? Já precisaram de capacidade extra?",
        "Que granulometria vocês precisam? Que tipo de material (PP, PE, PET)? Que cor?",
        "Como vem o material (fardo, bobina, solto)? Quem faz o frete?",
        "Precisam de MTR ou certificado de destinação (CDF)?",
    ]),
    ("3. Proposta", [
        "Opção 1: vocês mandam a apara e recebem o moído de volta (cobrança por kg).",
        "Opção 2 (se for gerador): compramos ou retiramos a apara.",
        "Oferecer teste grátis com um lote pequeno (ex.: 500 kg) para provar a qualidade.",
    ]),
    ("4. Fechamento", [
        "Posso enviar uma proposta por e-mail ou WhatsApp? Qual o melhor contato?",
        "Anotar na aba Ligar: nome, cargo, volume por mês e próximo passo, com data.",
    ]),
    ("Dicas", [
        "Ligue entre 9h e 11h30 ou entre 14h e 16h30. Peça pelo setor de Compras, Suprimentos, Produção ou Meio Ambiente.",
        "Nas fábricas (prioridade B), pergunte por 'resíduos' ou 'meio ambiente'. Nas recicladoras (prioridade A), por 'compras' ou 'produção'.",
        "Se não houver interesse, pergunte: 'Conhece alguém que gera ou compra esse tipo de apara?'",
    ]),
]
r = 3
for titulo, itens in roteiro:
    ws.cell(r, 1, titulo).font = Font(name=F, size=11, bold=True, color="FFFFFF")
    ws.cell(r, 1).fill = HEAD_FILL
    r += 1
    for it in itens:
        c = ws.cell(r, 1, "• " + it); c.font = BODY; c.alignment = WRAP
        r += 1
    r += 1
ws.column_dimensions["A"].width = 130

wb.save(OUT)
print(OUT, len(LIGAR), len(CONCORRENTES))
